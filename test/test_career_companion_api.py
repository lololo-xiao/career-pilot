from __future__ import annotations

from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.auth import AuthStore, AuthenticatedAccount, ProviderConnection
from app.main import app, get_store, require_current_account
from career_companion.persistence import clear_factory_cache


def _account(user_id: str) -> AuthenticatedAccount:
    return AuthenticatedAccount(
        user_id=user_id,
        email=f"{user_id}@example.test",
        display_name=user_id,
        identity_method="local",
        active_provider="api_key",
        provider_connection=ProviderConnection(
            provider="api_key",
            credential=b"sk-test-key-with-enough-characters",
        ),
    )


@pytest.fixture
def client(tmp_path, monkeypatch):
    clear_factory_cache()
    monkeypatch.setenv("CAREER_COMPANION_HOME", str(tmp_path / "companion"))
    store = AuthStore(tmp_path / "auth.db", "s" * 48)
    app.dependency_overrides[get_store] = lambda: store
    app.dependency_overrides[require_current_account] = lambda: _account("account-a")
    with TestClient(app, raise_server_exceptions=False) as test_client:
        yield test_client
    app.dependency_overrides.clear()
    clear_factory_cache()


def _job_payload() -> dict:
    return {
        "spec": {
            "title": "AI Engineer",
            "company": "Example GmbH",
            "locations": ["Berlin, Germany"],
            "description": "We require Python, retrieval, and evaluation experience.",
            "requirements": ["Python", "retrieval", "evaluation"],
            "source_type": "manual",
        },
        "canonical_url": "https://jobs.example.test/ai-engineer?utm_source=newsletter",
    }


def test_versioned_companion_routes_require_provider_connection(client) -> None:
    app.dependency_overrides.pop(require_current_account)

    response = client.get("/api/v1/jobs")

    assert response.status_code == 409
    assert response.json() == {
        "detail": "Choose an AI connection in Settings before using CareerPilot"
    }


def test_jobs_are_deduplicated_and_isolated_by_account(client) -> None:
    created = client.post("/api/v1/jobs", json=_job_payload())
    duplicate = client.post("/api/v1/jobs", json=_job_payload())

    assert created.status_code == 200
    assert created.json()["created"] is True
    assert created.json()["canonical_url"] == "https://jobs.example.test/ai-engineer"
    assert duplicate.status_code == 200
    assert duplicate.json()["created"] is False
    assert len(client.get("/api/v1/jobs").json()) == 1

    app.dependency_overrides[require_current_account] = lambda: _account("account-b")
    assert client.get("/api/v1/jobs").json() == []


def test_generic_revision_api_rejects_memory_before_any_mutation(client) -> None:
    response = client.post(
        "/api/v1/revisions",
        json={
            "kind": "memory",
            "name": "arbitrary-memory",
            "content": {"value": "unsafe"},
            "diff": "bypass",
            "author": "agent",
            "source_session": "session-1",
        },
    )

    assert response.status_code == 422
    assert client.get("/api/v1/revisions").json() == []
    assert not any(
        event["event_type"].startswith("revision.")
        for event in client.get("/api/v1/audit").json()
    )


def test_job_metadata_round_trips_for_queue_filters(client) -> None:
    payload = _job_payload()
    payload["spec"].update(
        {
            "posted_date": "2026-07-15",
            "workplace_type": "hybrid",
            "company_size": "51-200",
        }
    )

    created = client.post("/api/v1/jobs", json=payload)

    assert created.status_code == 200
    assert created.json()["spec"]["posted_date"] == "2026-07-15"
    assert created.json()["spec"]["workplace_type"] == "hybrid"
    assert created.json()["spec"]["company_size"] == "51-200"


def test_profile_api_rejects_unsupported_verified_claim(client) -> None:
    response = client.put(
        "/api/v1/onboarding/profile",
        json={
            "display_name": "Candidate",
            "claims": [
                {
                    "key": "skill",
                    "value": "Kubernetes",
                    "status": "verified",
                    "evidence": [],
                }
            ],
        },
    )

    assert response.status_code == 422
    assert "must include evidence" in response.json()["detail"]


def test_profile_upload_keeps_the_original_filename_in_citations(client) -> None:
    payload = BytesIO()
    document = Document()
    document.add_paragraph("Education")
    document.add_paragraph("MSc Computer Science, Example University")
    document.save(payload)

    response = client.post(
        "/api/v1/onboarding/import",
        files={
            "file": (
                "Ada Candidate Resume.docx",
                payload.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert response.status_code == 200
    claim = response.json()["profile"]["claims"][0]
    assert claim["evidence"][0]["source_name"] == "Ada Candidate Resume.docx"
    assert not claim["evidence"][0]["source_name"].startswith("upload-")


def test_local_project_analysis_round_trips_with_profile(client, tmp_path) -> None:
    repository = tmp_path / "sample-project"
    repository.mkdir()
    (repository / "package.json").write_text(
        '{"dependencies":{"next":"16","react":"19"}}', encoding="utf-8"
    )
    (repository / "tsconfig.json").write_text("{}", encoding="utf-8")
    (repository / "app.tsx").write_text("export default function App() {}", encoding="utf-8")
    hidden_runtime = repository / ".runtime"
    hidden_runtime.mkdir()
    (hidden_runtime / "generated.c").write_text("generated", encoding="utf-8")

    analyzed = client.post(
        "/api/v1/onboarding/projects/analyze",
        json={
            "id": "project-1",
            "name": "Sample project",
            "local_path": str(repository),
            "technologies": [],
            "highlights": [],
        },
    )

    assert analyzed.status_code == 200
    analysis = analyzed.json()
    assert analysis["source"] == "local"
    assert analysis["file_count"] == 3
    assert {"Next.js", "React", "TypeScript"}.issubset(analysis["technologies"])
    assert analysis["improvement_suggestions"]
    assert "Sample project" in analysis["interview_questions"][0]

    saved = client.put(
        "/api/v1/onboarding/profile",
        json={
            "display_name": "Ada Candidate",
            "seniority": "senior",
            "projects": [
                {
                    "id": "project-1",
                    "name": "Sample project",
                    "local_path": str(repository),
                    "technologies": analysis["technologies"],
                    "highlights": ["Built the main retrieval workflow"],
                    "analysis": analysis,
                }
            ],
        },
    )

    assert saved.status_code == 200
    assert saved.json()["seniority"] == "senior"
    assert saved.json()["projects"][0]["analysis"]["file_count"] == 3


def test_application_state_api_requires_confirmed_submission(client) -> None:
    job_id = client.post("/api/v1/jobs", json=_job_payload()).json()["id"]
    started = client.post("/api/v1/applications", params={"job_id": job_id})
    assert started.status_code == 200
    application_id = started.json()["id"]
    repeated = client.post("/api/v1/applications", params={"job_id": job_id})
    assert repeated.status_code == 200
    assert repeated.json()["id"] == application_id
    assert len(client.get("/api/v1/applications").json()) == 1

    for status_name in ("scored", "approved", "tailoring", "ready"):
        response = client.post(
            f"/api/v1/applications/{application_id}/status",
            json={"status": status_name},
        )
        assert response.status_code == 200

    unconfirmed = client.post(
        f"/api/v1/applications/{application_id}/status",
        json={"status": "submitted"},
    )
    assert unconfirmed.status_code == 409
    assert "explicit user confirmation" in unconfirmed.json()["detail"]

    confirmed = client.post(
        f"/api/v1/applications/{application_id}/status",
        json={"status": "submitted", "confirmed_by_user": True},
    )
    assert confirmed.status_code == 200
    assert confirmed.json()["status"] == "submitted"
    assert confirmed.json()["submitted_at"] is not None


def test_user_can_record_a_detailed_application_outcome(client) -> None:
    job_id = client.post("/api/v1/jobs", json=_job_payload()).json()["id"]
    application_id = client.post(
        "/api/v1/applications", params={"job_id": job_id}
    ).json()["id"]

    unconfirmed = client.post(
        f"/api/v1/applications/{application_id}/status",
        json={"status": "oa_failed", "manual_override": True},
    )
    recorded = client.post(
        f"/api/v1/applications/{application_id}/status",
        json={
            "status": "oa_failed",
            "manual_override": True,
            "confirmed_by_user": True,
            "note": "Recorded in the local workspace",
        },
    )

    assert unconfirmed.status_code == 409
    assert recorded.status_code == 200
    assert recorded.json()["status"] == "oa_failed"
    assert recorded.json()["submitted_at"] is not None
    assert recorded.json()["next_action"] == "Record learnings from the assessment"
    assert recorded.json()["status_events"][-1]["to"] == "oa_failed"


@pytest.mark.parametrize("protected_status", ["form_previewed", "form_filled"])
def test_generic_status_api_rejects_dedicated_form_workflow_targets(
    client,
    protected_status,
) -> None:
    job_id = client.post("/api/v1/jobs", json=_job_payload()).json()["id"]
    application_id = client.post(
        "/api/v1/applications",
        params={"job_id": job_id},
    ).json()["id"]

    response = client.post(
        f"/api/v1/applications/{application_id}/status",
        json={"status": protected_status, "manual_override": True},
    )

    assert response.status_code == 409
    assert "dedicated workflows" in response.json()["detail"]


def test_job_queue_csv_import_accepts_aliases_and_reports_row_errors(client) -> None:
    csv_text = """company,role,location,job_url,posted_date,workplace,company_size,description
Northstar Labs,AI Engineer,"Berlin, Germany",https://jobs.example.com/northstar-ai,2026-07-15,hybrid,51-200,Build retrieval systems.
Broken Row,Data Engineer,Paris,https://jobs.example.com/broken,not-a-date,remote,11-50,Build pipelines.
"""

    response = client.post(
        "/api/v1/jobs/import",
        files={"file": ("jobs.csv", csv_text, "text/csv")},
    )

    assert response.status_code == 200
    assert response.json()["created_jobs"] == 1
    assert len(response.json()["errors"]) == 1
    jobs = client.get("/api/v1/jobs").json()
    assert jobs[0]["spec"]["workplace_type"] == "hybrid"
    assert jobs[0]["spec"]["company_size"] == "51-200"


def test_application_csv_import_creates_jobs_and_detailed_statuses(client) -> None:
    csv_text = """company,title,location,url,status,applied_date,workplace,company_size,notes
Northstar Labs,AI Engineer,"Berlin, Germany",https://jobs.example.com/northstar-ai,Applied,2026-07-10,hybrid,51-200,Applied manually
Pinecone Robotics,ML Engineer,"Munich, Germany",https://jobs.example.com/pinecone-ml,Interview 1 Fail,2026-07-08,onsite,201-500,Review system design
"""

    first = client.post(
        "/api/v1/applications/import",
        files={"file": ("applications.csv", csv_text, "text/csv")},
    )
    second = client.post(
        "/api/v1/applications/import",
        files={"file": ("applications.csv", csv_text, "text/csv")},
    )

    assert first.status_code == 200
    assert first.json()["created_jobs"] == 2
    assert first.json()["created_applications"] == 2
    assert second.status_code == 200
    assert second.json()["created_jobs"] == 0
    assert second.json()["created_applications"] == 0
    assert second.json()["updated_applications"] == 2
    applications = client.get("/api/v1/applications").json()
    assert {item["status"] for item in applications} == {
        "submitted",
        "interview_1_failed",
    }
    assert all(item["submitted_at"].startswith("2026-07") for item in applications)


@pytest.mark.parametrize("protected_status", ["form_previewed", "form_filled"])
def test_application_csv_import_rejects_dedicated_workflow_statuses(
    client,
    protected_status,
) -> None:
    csv_text = f"""company,title,url,status,notes
Boundary Labs,AI Engineer,https://jobs.example.com/{protected_status},{protected_status},Imported status
"""

    response = client.post(
        "/api/v1/applications/import",
        files={"file": ("applications.csv", csv_text, "text/csv")},
    )

    assert response.status_code == 200
    result = response.json()
    assert result["created_jobs"] == 0
    assert result["created_applications"] == 0
    assert result["errors"] == [
        f"Row 2: status '{protected_status}' requires a dedicated workflow and cannot be imported"
    ]
    assert client.get("/api/v1/applications").json() == []


def test_demo_csv_templates_import_end_to_end(client) -> None:
    template_dir = Path(__file__).parents[1] / "frontend" / "public" / "templates"

    with (template_dir / "job-queue-demo-template.csv").open("rb") as jobs_file:
        jobs_response = client.post(
            "/api/v1/jobs/import",
            files={"file": ("jobs.csv", jobs_file, "text/csv")},
        )
    with (template_dir / "applications-demo-template.csv").open("rb") as applications_file:
        applications_response = client.post(
            "/api/v1/applications/import",
            files={"file": ("applications.csv", applications_file, "text/csv")},
        )

    assert jobs_response.status_code == 200
    assert jobs_response.json()["created_jobs"] == 20
    assert jobs_response.json()["errors"] == []
    assert applications_response.status_code == 200
    assert applications_response.json()["created_applications"] == 20
    assert applications_response.json()["errors"] == []
    assert len(client.get("/api/v1/jobs").json()) == 20
    assert len(client.get("/api/v1/applications").json()) == 20


def test_integrated_api_uses_existing_provider_settings_only(client) -> None:
    duplicate_writer = client.post(
        "/api/v1/provider-auth/openai-api",
        data={"api_key": "sk-this-must-never-be-written-here"},
    )
    existing_settings = client.get("/settings/providers")

    assert duplicate_writer.status_code == 404
    assert existing_settings.status_code == 200


def test_default_model_routes_are_initialized_per_account(client) -> None:
    response = client.get("/api/v1/model-routes")

    assert response.status_code == 200
    routes = {route["name"]: route for route in response.json()}
    assert set(routes) == {
        "interactive",
        "research",
        "extraction",
        "tailoring",
        "evaluation",
        "memory-review",
        "compression",
        "cron",
    }
    assert routes["cron"]["scheduled"] is True
    assert routes["cron"]["fallback_policy"] == "none"
