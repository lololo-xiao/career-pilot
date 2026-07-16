from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document

from app.ingestion import (
    CVParseError,
    CVTooLargeError,
    MAX_DOCX_UNCOMPRESSED_BYTES,
    MAX_PDF_PAGES,
    MAX_UPLOAD_BYTES,
    UnsupportedCVTypeError,
    parse_cv_document,
)


def make_text_pdf(text: str) -> bytes:
    content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>"
        ),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length "
        + str(len(content)).encode()
        + b" >>\nstream\n"
        + content
        + b"\nendstream",
    ]
    pdf = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf.extend(f"{number} 0 obj\n".encode())
        pdf.extend(body)
        pdf.extend(b"\nendobj\n")

    xref_offset = len(pdf)
    pdf.extend(f"xref\n0 {len(objects) + 1}\n".encode())
    pdf.extend(b"0000000000 65535 f \n")
    for offset in offsets:
        pdf.extend(f"{offset:010d} 00000 n \n".encode())
    pdf.extend(
        (
            f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref_offset}\n%%EOF\n"
        ).encode()
    )
    return bytes(pdf)


def test_parse_utf8_text_cv_and_sanitize_filename() -> None:
    result = parse_cv_document(
        "../candidate.txt",
        b"Python engineer\r\n\r\nBuilt RAG systems with FastAPI.",
    )

    assert result.filename == "candidate.txt"
    assert result.file_type == "text"
    assert result.text == "Python engineer\n\nBuilt RAG systems with FastAPI."
    assert result.character_count == len(result.text)
    assert result.page_count is None


def test_parse_docx_includes_paragraphs_and_tables() -> None:
    document = Document()
    document.add_heading("AI Engineer", level=1)
    document.add_paragraph("Built Python and RAG services.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Skill"
    table.cell(0, 1).text = "FastAPI"
    buffer = BytesIO()
    document.save(buffer)

    result = parse_cv_document("candidate.docx", buffer.getvalue())

    assert result.file_type == "docx"
    assert "AI Engineer" in result.text
    assert "Built Python and RAG services." in result.text
    assert "Skill\tFastAPI" in result.text


def test_parse_digitally_generated_pdf() -> None:
    result = parse_cv_document(
        "candidate.pdf",
        make_text_pdf("Python engineer with RAG experience."),
    )

    assert result.file_type == "pdf"
    assert result.page_count == 1
    assert "Python engineer with RAG experience." in result.text


def test_rejects_image_only_or_empty_pdf() -> None:
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    buffer = BytesIO()
    writer.write(buffer)

    with pytest.raises(CVParseError, match="No usable text"):
        parse_cv_document("scan.pdf", buffer.getvalue())


def test_rejects_unsupported_and_oversized_files() -> None:
    with pytest.raises(UnsupportedCVTypeError, match="PDF, DOCX, and TXT"):
        parse_cv_document("candidate.rtf", b"Some candidate content")

    with pytest.raises(CVTooLargeError, match="5 MB"):
        parse_cv_document("candidate.txt", b"x" * (MAX_UPLOAD_BYTES + 1))


def test_rejects_docx_archive_expansion_bomb() -> None:
    buffer = BytesIO()
    with ZipFile(buffer, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types />")
        archive.writestr("word/document.xml", "<document />")
        archive.writestr(
            "word/media/oversized.bin",
            b"0" * (MAX_DOCX_UNCOMPRESSED_BYTES + 1),
        )

    assert len(buffer.getvalue()) < MAX_UPLOAD_BYTES
    with pytest.raises(CVTooLargeError, match="expanded DOCX"):
        parse_cv_document("candidate.docx", buffer.getvalue())


def test_rejects_pdf_with_too_many_pages() -> None:
    from pypdf import PdfWriter

    writer = PdfWriter()
    for _ in range(MAX_PDF_PAGES + 1):
        writer.add_blank_page(width=612, height=792)
    buffer = BytesIO()
    writer.write(buffer)

    with pytest.raises(CVTooLargeError, match=f"{MAX_PDF_PAGES} pages"):
        parse_cv_document("long-cv.pdf", buffer.getvalue())
