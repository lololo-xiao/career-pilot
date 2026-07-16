import re
from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from app.schemas import CVFileType, ParsedCVResponse


MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_EXTRACTED_CHARS = 50_000
MAX_PDF_PAGES = 30
MAX_DOCX_ENTRIES = 500
MAX_DOCX_UNCOMPRESSED_BYTES = 20 * 1024 * 1024

_MULTIPLE_BLANK_LINES = re.compile(r"\n{3,}")
_TRAILING_WHITESPACE = re.compile(r"[ \t]+\n")


class CVIngestionError(ValueError):
    pass


class CVTooLargeError(CVIngestionError):
    pass


class UnsupportedCVTypeError(CVIngestionError):
    pass


class CVParseError(CVIngestionError):
    pass


def _normalize_text(text: str) -> str:
    normalized = text.replace("\r\n", "\n").replace("\r", "\n")
    normalized = _TRAILING_WHITESPACE.sub("\n", normalized)
    normalized = _MULTIPLE_BLANK_LINES.sub("\n\n", normalized).strip()
    if len(normalized) < 10:
        raise CVParseError(
            "No usable text was found. Scanned PDFs require OCR, which is not supported."
        )
    if len(normalized) > MAX_EXTRACTED_CHARS:
        raise CVTooLargeError(
            f"Extracted CV text exceeds {MAX_EXTRACTED_CHARS:,} characters"
        )
    return normalized


def _parse_pdf(data: bytes) -> tuple[str, int]:
    if not data.startswith(b"%PDF-"):
        raise CVParseError("The uploaded file is not a valid PDF")
    try:
        reader = PdfReader(BytesIO(data), strict=False)
        page_count = len(reader.pages)
        if page_count == 0:
            raise CVParseError("The PDF contains no pages")
        if page_count > MAX_PDF_PAGES:
            raise CVTooLargeError(
                f"PDFs are limited to {MAX_PDF_PAGES} pages for the demo"
            )
        pages = [page.extract_text() or "" for page in reader.pages]
    except (PdfReadError, OSError, ValueError) as exc:
        if isinstance(exc, CVIngestionError):
            raise
        raise CVParseError("The PDF could not be parsed") from exc
    return _normalize_text("\n\n".join(pages)), page_count


def _validate_docx_archive(data: bytes) -> None:
    try:
        with ZipFile(BytesIO(data)) as archive:
            entries = archive.infolist()
            names = {entry.filename for entry in entries}
            if "[Content_Types].xml" not in names or "word/document.xml" not in names:
                raise CVParseError("The uploaded file is not a valid DOCX document")
            if len(entries) > MAX_DOCX_ENTRIES:
                raise CVTooLargeError("The DOCX archive contains too many files")
            total_size = sum(entry.file_size for entry in entries)
            if total_size > MAX_DOCX_UNCOMPRESSED_BYTES:
                raise CVTooLargeError("The expanded DOCX document is too large")
    except BadZipFile as exc:
        raise CVParseError("The uploaded file is not a valid DOCX document") from exc


def _parse_docx(data: bytes) -> str:
    _validate_docx_archive(data)
    try:
        document = Document(BytesIO(data))
    except (PackageNotFoundError, ValueError, KeyError) as exc:
        raise CVParseError("The DOCX document could not be parsed") from exc

    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            row_text = "\t".join(
                cell.text.strip() for cell in row.cells if cell.text.strip()
            )
            if row_text:
                parts.append(row_text)
    return _normalize_text("\n".join(parts))


def _parse_text(data: bytes) -> str:
    try:
        return _normalize_text(data.decode("utf-8-sig"))
    except UnicodeDecodeError as exc:
        raise CVParseError("Text CVs must use UTF-8 encoding") from exc


def parse_cv_document(filename: str, data: bytes) -> ParsedCVResponse:
    """Extract bounded plain text from a supported in-memory CV document."""

    safe_filename = Path(filename).name
    if not safe_filename:
        raise UnsupportedCVTypeError("The uploaded CV needs a filename")
    if not data:
        raise CVParseError("The uploaded CV is empty")
    if len(data) > MAX_UPLOAD_BYTES:
        raise CVTooLargeError("CV uploads are limited to 5 MB")

    extension = Path(safe_filename).suffix.casefold()
    file_type: CVFileType
    page_count: int | None = None
    if extension == ".pdf":
        file_type = "pdf"
        text, page_count = _parse_pdf(data)
    elif extension == ".docx":
        file_type = "docx"
        text = _parse_docx(data)
    elif extension in {".txt", ".md"}:
        file_type = "text"
        text = _parse_text(data)
    else:
        raise UnsupportedCVTypeError("Supported CV formats are PDF, DOCX, and TXT")

    return ParsedCVResponse(
        filename=safe_filename,
        file_type=file_type,
        text=text,
        character_count=len(text),
        page_count=page_count,
    )
